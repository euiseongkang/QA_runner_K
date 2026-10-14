#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""TC 엑셀 파서 - 프로그램(qa_runner_k_gui.py)과 대시보드(dashboard_server.py)가 함께 쓴다.

두 곳에서 같은 엑셀을 읽는데 파싱 규칙이 따로 놀면 "프로그램에서는 되는데 대시보드에서는
안 되는" 상황이 생기므로, 규칙을 이 한 곳에만 둔다. [NEW v0.5.0]

지원 포맷 (강의성님이 확정한 표준):
  시트 "개요"        - 문서 정보 (파싱하지 않음)
  시트 "테스트케이스" - No / 테스트 항목 / 사전조건 / 테스트 절차 / 예상 결과 / 우선순위 /
                       결과(Pass/Fail) / 비고

실제 파일에 맞춰 들어간 규칙:
  - 헤더는 공백을 무시하고 비교 ("테스트 항목" = "테스트항목")
  - 실행에 꼭 필요한 건 테스트 항목 / 테스트 절차 / 예상 결과 3개. 나머지는 비어도 통과
  - No가 비어 있는 행이 실제로 많아서, 비면 엑셀 행 번호(r<번호>)로 대체
  - 완전히 빈 행은 조용히 건너뜀 (작성용 빈 줄이 대부분인 파일이 있음)
"""

import re

REQUIRED_COLUMNS = ["테스트 항목", "테스트 절차", "예상 결과"]


def norm_header(name) -> str:
    """헤더/컬럼명 비교용 정규화 (공백 제거)."""
    return re.sub(r"\s+", "", str(name or ""))


def clean_cell(value) -> str:
    """셀 값의 앞뒤 공백/중복 공백 정리. 줄바꿈은 살려야 하는 컬럼이 있어 여기서는 쓰지 않음.

    [v0.16.0] 숫자 셀이 실수로 넘어오는 경우를 정리한다. 구글 시트를 거치면 No 컬럼의 1이
    1.0으로 와서 TC 번호가 "1.0"으로 보이는 문제가 있었다."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"[ \t]+", " ", str(value)).strip()


class TCExcelError(ValueError):
    """포맷이 달라서 읽을 수 없을 때. 메시지를 그대로 사용자에게 보여준다."""


# ============================================================
# TC 시트 고르기                                                [NEW v0.14.0]
# ============================================================
# 처음에는 "테스트케이스" 시트 하나만 읽었는데, 실제로는 화면별로 시트를 나눠 쓰고
# 이름도 "TC_환자관리"처럼 자유롭게 붙인다. 그래서 이름에 TC(또는 테스트케이스)가 들어간
# 시트를 전부 읽고, 시트 이름을 TC의 구분값으로 달아둔다.
SHEET_NAME = "테스트케이스"            # 기본/예시 이름 (하위호환)
SHEET_KEYWORDS = ("TC", "테스트케이스")


def is_tc_sheet(name) -> bool:
    """시트 이름에 TC 또는 테스트케이스가 들어가면 TC 시트로 본다 (대소문자·공백 무시).

    "개요"처럼 TC가 아닌 시트를 걸러내는 게 목적이라 조건을 느슨하게 둔다.
    필수 컬럼이 없으면 어차피 읽는 단계에서 걸러진다."""
    key = norm_header(name).upper()
    if not key:
        return False
    return any(k.upper() in key for k in SHEET_KEYWORDS)


def pick_sheets(sheetnames):
    """워크북의 시트 이름 목록에서 TC 시트만 원래 순서대로 고른다."""
    return [n for n in (sheetnames or []) if is_tc_sheet(n)]


def _parse_one_sheet(ws, sheet_name, strict=True):
    """시트 하나를 읽어 (tcs, warnings). 필수 컬럼이 없으면 그 시트만 건너뛴다.

    여러 시트를 읽게 되면서, 시트 하나가 잘못됐다고 파일 전체를 실패시키면
    나머지 멀쩡한 시트까지 못 쓰게 되므로 경고만 남기고 넘어간다.
    strict=False면 이름으로 지목된 시트가 아니라 훑어보는 중이므로 경고도 남기지 않는다."""
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return [], ([f'"{sheet_name}" 시트가 비어 있어 건너뜀'] if strict else [])

    col = {}
    for idx, name in enumerate(header_row):
        key = norm_header(name)
        if key and key not in col:
            col[key] = idx

    missing = [c for c in REQUIRED_COLUMNS if norm_header(c) not in col]
    if missing:
        if not strict:
            return [], []      # TC 시트가 아닌 걸로 보고 조용히 넘어간다 (예: "개요")
        return [], [f'"{sheet_name}" 시트 건너뜀 - 필수 컬럼 없음: ' + ", ".join(missing)]

    def cell(row, name, keep_newlines=False):
        idx = col.get(norm_header(name))
        if idx is None or idx >= len(row) or row[idx] is None:
            return ""
        return str(row[idx]) if keep_newlines else clean_cell(row[idx])

    tcs, warnings = [], []
    for row_idx, row in enumerate(rows, start=2):
        title = cell(row, "테스트 항목")
        steps = cell(row, "테스트 절차", keep_newlines=True)
        expected = cell(row, "예상 결과")

        if not title and not steps.strip() and not expected:
            continue  # 완전히 빈 행

        if not title or not steps.strip() or not expected:
            missing_here = [
                n for n, v in (("테스트 항목", title), ("테스트 절차", steps.strip()),
                               ("예상 결과", expected)) if not v
            ]
            warnings.append(f'"{sheet_name}" {row_idx}행 건너뜀 - 빈 칸: ' + ", ".join(missing_here))
            continue

        no = cell(row, "No")
        tcs.append({
            "no": no or f"r{row_idx}",
            "title": title,
            "precondition": cell(row, "사전조건"),
            "steps": steps,
            "expected": expected,
            "priority": cell(row, "우선순위"),
            "note": cell(row, "비고"),
            "sheet": sheet_name,          # [NEW v0.14.0] 시트별 관리를 위한 구분값
            "row": row_idx,
        })
    return tcs, warnings


# ============================================================
# 구글 스프레드시트에서 바로 가져오기                          [NEW v0.13.0]
# ============================================================
# 구글 시트는 파일이 아니라 웹 문서라 그대로는 못 읽는다. 대신 구글이 제공하는
# "xlsx로 내보내기" 주소로 받아서 평소와 똑같이 파싱한다. CSV가 아니라 xlsx로 받는 이유는,
# CSV는 시트 하나만 나와서 "테스트케이스" 시트 이름이 사라지기 때문.
GSHEET_HOSTS = ("docs.google.com",)
GSHEET_MAX_BYTES = 10 * 1024 * 1024


def is_gsheet_url(url) -> bool:
    return "docs.google.com/spreadsheets" in str(url or "")


def gsheet_export_url(url) -> str:
    """구글 시트 주소를 xlsx 내보내기 주소로 바꾼다.

    받아들이는 형태
      .../spreadsheets/d/<ID>/edit#gid=0        (일반 공유 링크)
      .../spreadsheets/d/<ID>                    (짧은 형태)
      .../spreadsheets/d/e/<PUBID>/pubhtml       (웹에 게시한 링크)
    """
    import urllib.parse as _up

    url = str(url or "").strip()
    if not url:
        raise TCExcelError("구글 시트 주소를 입력해주세요")

    parsed = _up.urlparse(url if "://" in url else "https://" + url)
    # 주소를 그대로 받아 서버가 대신 요청하는 구조라, 사내망 주소 같은 걸 넣어
    # 서버를 심부름꾼으로 쓰지 못하게 호스트를 구글로 못박는다 (SSRF 방지).
    if parsed.scheme not in ("http", "https") or parsed.hostname not in GSHEET_HOSTS:
        raise TCExcelError("구글 스프레드시트 주소(docs.google.com)만 가져올 수 있습니다")

    m = re.search(r"/spreadsheets/d/e/([A-Za-z0-9\-_]+)", parsed.path)
    if m:                      # 웹에 게시한 문서
        return f"https://docs.google.com/spreadsheets/d/e/{m.group(1)}/pub?output=xlsx"
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9\-_]+)", parsed.path)
    if m:
        return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx"
    raise TCExcelError("구글 시트 주소에서 문서 ID를 찾지 못했습니다. 주소창의 링크를 그대로 붙여넣어 주세요")


def _name_from_disposition(value):
    """Content-Disposition 헤더에서 파일 이름을 꺼낸다. [NEW v0.32.0]

    구글은 보통 두 가지를 같이 준다.
        filename="LabConnect TC.xlsx"
        filename*=UTF-8''LabConnect%20TC.xlsx
    한글 제목은 앞쪽이 깨져 오는 경우가 있어 RFC 5987 형태(filename*)를 먼저 본다.
    못 읽으면 빈 문자열을 돌려주고, 부르는 쪽에서 기본 이름으로 대체한다.
    """
    import urllib.parse as _up

    text = str(value or "")
    if not text:
        return ""
    m = re.search(r"filename\*\s*=\s*([^']*)'[^']*'([^;]+)", text, re.IGNORECASE)
    if m:
        charset = (m.group(1) or "utf-8").strip() or "utf-8"
        raw = m.group(2).strip().strip('"')
        try:
            name = _up.unquote(raw, encoding=charset, errors="replace")
        except (LookupError, ValueError):
            name = _up.unquote(raw)
        if name.strip():
            return name.strip()
    m = re.search(r'filename\s*=\s*"([^"]*)"', text, re.IGNORECASE)
    if not m:
        m = re.search(r"filename\s*=\s*([^;]+)", text, re.IGNORECASE)
    return m.group(1).strip().strip('"') if m else ""


def strip_excel_ext(name):
    """파일 이름 끝의 엑셀 확장자를 떼어낸다. [NEW v0.32.0]

    구글 시트는 '<시트 이름>.xlsx' 로 내려오므로, 시트 이름만 남기려면 떼야 한다.
    """
    text = str(name or "").strip()
    for ext in (".xlsx", ".xlsm", ".xls"):
        if text.lower().endswith(ext):
            return text[: -len(ext)].strip()
    return text


def fetch_gsheet_named(url, timeout=20):
    """구글 시트를 (xlsx 바이트, 시트 이름) 으로 받아온다. [NEW v0.32.0]

    시트 이름은 구글이 내려주는 Content-Disposition 에서 꺼낸다. 헤더가 없거나 못 읽으면
    빈 문자열을 돌려주므로, 부르는 쪽에서 기본값으로 대체하면 된다.
    **이름을 못 읽는다고 가져오기 자체가 실패하지는 않는다.**
    """
    import io as _io
    import urllib.request as _ur
    import urllib.error as _ue

    export = gsheet_export_url(url)
    req = _ur.Request(export, headers={"User-Agent": "QA_runner_K"})
    try:
        with _ur.urlopen(req, timeout=timeout) as resp:
            ctype = (resp.headers.get("Content-Type") or "").lower()
            disposition = resp.headers.get("Content-Disposition") or ""
            data = resp.read(GSHEET_MAX_BYTES + 1)
    except _ue.HTTPError as e:
        if e.code in (401, 403):
            raise TCExcelError("시트를 열 권한이 없습니다. 구글 시트에서 [공유] → "
                               "'링크가 있는 모든 사용자'(뷰어)로 바꾼 뒤 다시 시도해주세요")
        if e.code == 404:
            raise TCExcelError("시트를 찾을 수 없습니다. 주소가 맞는지 확인해주세요")
        raise TCExcelError(f"구글 시트를 받지 못했습니다 (HTTP {e.code})")
    except Exception as e:
        raise TCExcelError(f"구글 시트에 연결하지 못했습니다: {str(e)[:120]}")

    if len(data) > GSHEET_MAX_BYTES:
        raise TCExcelError("시트가 너무 큽니다 (10MB 초과)")
    # 비공개 시트는 오류 대신 로그인 페이지(HTML)를 200으로 돌려준다. xlsx는 'PK'로 시작한다.
    if not data.startswith(b"PK"):
        if "html" in ctype:
            raise TCExcelError("시트가 공개되어 있지 않습니다. 구글 시트에서 [공유] → "
                               "'링크가 있는 모든 사용자'(뷰어)로 바꾼 뒤 다시 시도해주세요")
        raise TCExcelError("구글 시트에서 받은 내용이 엑셀 형식이 아닙니다")

    name = ""
    try:
        name = strip_excel_ext(_name_from_disposition(disposition))
    except Exception:
        name = ""       # 이름은 덤이다. 못 읽어도 TC 가져오기는 그대로 진행한다.
    return _io.BytesIO(data), name


def fetch_gsheet(url, timeout=20):
    """구글 시트를 xlsx 바이트로 받아온다. 실패 사유는 사람이 바로 고칠 수 있게 풀어서 알린다.

    [v0.32.0] 본체는 fetch_gsheet_named 로 옮겼다. 이름이 필요 없는 기존 호출을 위해 남겨둔다.
    """
    stream, _name = fetch_gsheet_named(url, timeout=timeout)
    return stream


def parse_gsheet(url, timeout=20):
    """구글 시트 주소 -> (tcs, warnings). 파싱은 엑셀 업로드와 완전히 같은 경로를 탄다."""
    return parse_tc_excel(fetch_gsheet(url, timeout=timeout))


def parse_gsheet_named(url, timeout=20):
    """구글 시트 주소 -> (tcs, warnings, 시트 이름). [NEW v0.32.0]

    세션 이름을 시트 이름으로 달기 위해 추가했다. 파싱 경로는 parse_gsheet 와 완전히 같다.
    """
    stream, name = fetch_gsheet_named(url, timeout=timeout)
    tcs, warnings = parse_tc_excel(stream)
    return tcs, warnings, name


def parse_tc_excel(source):
    """TC 엑셀을 읽어 (tcs, warnings) 반환.

    source: 파일 경로(str) 또는 file-like 객체 (Flask 업로드 스트림도 그대로 사용 가능)
    tcs   : [{"no","title","precondition","steps","expected","priority","note","row"}, ...]
            - steps는 줄바꿈을 그대로 보존한다 (번호 매긴 절차가 줄 단위로 해석되므로)
    warnings: 건너뛴 행 등 사람에게 알려줄 메모 리스트
    """
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - 빌드 환경에서는 항상 포함됨
        raise TCExcelError("엑셀을 읽는 데 필요한 openpyxl이 설치되어 있지 않습니다")

    try:
        wb = load_workbook(source, data_only=True, read_only=True)
    except Exception as e:
        raise TCExcelError(f"엑셀 파일을 열 수 없습니다: {str(e)[:120]}")

    try:
        # [v0.16.0] 시트 이름은 힌트일 뿐이고, 이름에 TC가 없어도 컬럼만 맞으면 읽는다.
        # "환자 관리"처럼 TC를 안 붙인 이름을 쓰시는 경우가 있어서, 이름으로 막지 않는다.
        named = pick_sheets(wb.sheetnames)
        targets = named or list(wb.sheetnames)

        tcs, warnings = [], []
        for sheet_name in targets:
            sheet_tcs, sheet_warns = _parse_one_sheet(
                wb[sheet_name], sheet_name, strict=bool(named))
            tcs.extend(sheet_tcs)
            warnings.extend(sheet_warns)

        if not tcs and not named:
            raise TCExcelError(
                "TC를 읽을 수 있는 시트가 없습니다. 시트 1행 헤더에 "
                "'테스트 항목 · 테스트 절차 · 예상 결과'가 있어야 합니다. "
                f'(이 파일의 시트: {", ".join(wb.sheetnames) or "없음"})'
            )
        if not tcs:
            warnings.append("읽을 수 있는 TC가 없습니다 (데이터 행이 비어 있는지 확인해주세요)")
        return tcs, warnings
    finally:
        try:
            wb.close()
        except Exception:
            pass


# ============================================================
# TC 엑셀로 내보내기                                            [NEW v0.36.0]
# ============================================================
# 읽기와 쓰기를 같은 파일에 두는 이유는 이 모듈이 존재하는 이유와 같다 - 포맷 규칙이
# 두 곳으로 갈라지면 "내보낸 파일을 자기가 못 읽는" 상황이 생긴다.
# 그래서 내보내기의 목표는 "보기 좋은 엑셀"이 아니라 **parse_tc_excel 이 그대로 다시
# 읽을 수 있는 파일**이다. 아래 규칙은 전부 그 왕복을 위해 있다.
#
# 시트 이름을 TC 로 알아볼 수 있게 맞추는 이유 (중요):
#   parse_tc_excel 은 이름에 TC/테스트케이스가 든 시트가 **하나라도** 있으면
#   그 시트들만 읽는다. 그래서 한 세션에 "TC_검사"와 "환자관리"가 섞여 있으면
#   다시 가져올 때 "환자관리" 쪽 TC가 통째로 사라진다.
#   이 사고를 막으려고, 내보낼 때 TC로 안 보이는 이름에는 TC_ 를 붙인다.
#   이름이 바뀌는 건 눈에 보이지만, TC가 조용히 없어지는 것보다 낫다.

EXPORT_COLUMNS = ["No", "테스트 항목", "사전조건", "테스트 절차", "예상 결과",
                  "우선순위", "결과(Pass/Fail)", "비고"]
OVERVIEW_SHEET = "개요"
DEFAULT_TC_SHEET = "테스트케이스"

# 엑셀 시트 이름에 못 쓰는 문자. 넣으면 openpyxl 이 아니라 엑셀이 파일을 거부한다.
_SHEET_BAD = r'[]:*?/\\'
_SHEET_MAX = 31

# 윈도우 파일명에 못 쓰는 문자.
_FILE_BAD = r'\/:*?"<>|'
_FILE_MAX = 120


def safe_sheet_title(name, used=None):
    """시트 이름을 엑셀이 받아주고 parse_tc_excel 이 TC 시트로 알아보는 형태로 다듬는다.

    used 에 이미 쓴 이름을 넣어주면 중복을 피해 숫자를 붙인다(엑셀은 대소문자를 구분하지 않는다).
    """
    used = used if used is not None else set()
    title = str(name or "").strip()
    for ch in _SHEET_BAD:
        title = title.replace(ch, " ")
    title = re.sub(r"\s+", " ", title).strip().strip("'")
    if not title:
        title = DEFAULT_TC_SHEET
    if not is_tc_sheet(title):
        title = "TC_" + title           # 다시 읽힐 수 있게. 위 주석 참고
    title = title[:_SHEET_MAX].strip() or DEFAULT_TC_SHEET

    base, n = title, 2
    while title.lower() in used:
        suffix = f"_{n}"
        title = base[:_SHEET_MAX - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def safe_file_part(name):
    """파일명에 넣어도 되는 형태로 다듬는다. 빈 문자열이 나올 수 있다(그때는 호출한 쪽이 뺀다)."""
    part = str(name or "")
    for ch in _FILE_BAD:
        part = part.replace(ch, "_")
    part = "".join(c for c in part if ord(c) >= 32)
    part = re.sub(r"\s+", " ", part).strip().strip(".")
    return part


def export_filename(label="", today=None):
    """내보내기 파일명: TC_<날짜>_<세션 이름>.xlsx

    세션 이름을 붙이는 건 같은 날 여러 세션을 받을 때 브라우저가 "(1)"을 붙여
    어느 세션인지 알 수 없게 되는 걸 막기 위한 것이다.
    """
    import time as _time
    stamp = today or _time.strftime("%Y-%m-%d")
    part = safe_file_part(label)
    name = f"TC_{stamp}_{part}" if part else f"TC_{stamp}"
    return name[:_FILE_MAX].strip().strip(".") + ".xlsx"


def _overview_rows(session, tcs, sheet_renames=None):
    """개요 시트에 넣을 (항목, 내용) 줄들. 표준 포맷의 항목명을 그대로 쓴다."""
    import time as _time
    session = session or {}

    def when(ts):
        try:
            return _time.strftime("%Y-%m-%d %H:%M", _time.localtime(float(ts)))
        except (TypeError, ValueError):
            return ""

    sheets = sorted({str(t.get("sheet") or "") for t in tcs} - {""})
    rows = [
        ("문서명", str(session.get("label") or "QA 테스트케이스")),
        ("대상 화면", ", ".join(sheets) or "-"),
        ("작성일", when(session.get("created_at")) or "-"),
        ("작성 근거",
         "QA_runner_K 대시보드에 등록된 TC를 그대로 내보낸 파일입니다. "
         "이 파일은 대시보드의 [TC 엑셀 올리기]로 다시 가져올 수 있습니다."),
        ("범위 안내",
         "결과(Pass/Fail) 칸은 비어 있습니다. 실행 결과는 대시보드의 실행 내역에서 봅니다. "
         "실행 포함/제외 상태는 이 파일에 담기지 않으므로, 다시 가져오면 전부 '포함'으로 들어옵니다."),
        ("", ""),
        ("세션 ID", str(session.get("session_id") or "-")),
        ("출처", str(session.get("source") or "-")),
        ("불러온 시각", when(session.get("created_at")) or "-"),
        ("내보낸 시각", _time.strftime("%Y-%m-%d %H:%M")),
        ("TC 건수", f"{len(tcs)}건"),
    ]
    if sheet_renames:
        rows.append(
            ("시트 이름 변경",
             "다시 가져올 때 TC가 빠지지 않도록 아래 시트 이름 앞에 TC_ 를 붙였습니다: "
             + ", ".join(f"{a} -> {b}" for a, b in sheet_renames)))
    return rows


def build_tc_workbook(tcs, session=None):
    """TC 목록을 표준 포맷(개요 + TC 시트들) 엑셀로 만들어 BytesIO 로 돌려준다.

    tcs: list_custom_tcs() 가 주는 dict 목록 (title/steps/expected/... 키를 쓴다)
    원본 시트별로 시트를 나누고, 세션 안에서의 순서를 그대로 유지한다.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:  # pragma: no cover - 빌드 환경에는 항상 포함된다
        raise TCExcelError("엑셀을 만드는 데 필요한 openpyxl이 설치되어 있지 않습니다")

    tcs = list(tcs or [])

    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="EEEEEE")
    top_wrap = Alignment(vertical="top", wrap_text=True)
    top_plain = Alignment(vertical="top")

    # ---- 시트별로 묶기 (세션 안 순서 유지) ----
    groups, order = {}, []
    for tc in tcs:
        key = str(tc.get("sheet") or "")
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(tc)
    if not order:                      # TC가 하나도 없어도 빈 표는 만들어준다
        groups[""], order = [], [""]

    used = {OVERVIEW_SHEET.lower()}
    titles, renames = {}, []
    for key in order:
        title = safe_sheet_title(key or DEFAULT_TC_SHEET, used)
        titles[key] = title
        if key and title != key:
            renames.append((key, title))

    wb = Workbook()
    wb.remove(wb.active)

    # ---- 개요 ----
    ov = wb.create_sheet(OVERVIEW_SHEET)
    ov.append(["항목", "내용"])
    for cell in ov[1]:
        cell.font, cell.fill = head_font, head_fill
    for item, value in _overview_rows(session, tcs, renames):
        ov.append([item, value])
    for row in ov.iter_rows(min_row=2):
        row[0].font = head_font
        row[0].alignment = top_plain
        row[1].alignment = top_wrap
    ov.column_dimensions["A"].width = 16
    ov.column_dimensions["B"].width = 100
    ov.freeze_panes = "A2"

    # ---- TC 시트들 ----
    widths = [8, 34, 24, 52, 44, 10, 14, 24]
    for key in order:
        ws = wb.create_sheet(titles[key])
        ws.append(EXPORT_COLUMNS)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = top_plain
        for tc in groups[key]:
            ws.append([
                str(tc.get("tc_no") or ""),
                str(tc.get("title") or ""),
                str(tc.get("precondition") or ""),
                str(tc.get("steps") or ""),
                str(tc.get("expected") or ""),
                str(tc.get("priority") or ""),
                "",                                  # 결과(Pass/Fail) - 실행 후 채우는 칸
                str(tc.get("note") or ""),
            ])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = top_wrap
            # No 는 숫자로 해석돼 "1"이 "1.0"으로 돌아오는 일이 있어 글자로 고정한다
            row[0].number_format = "@"
            row[0].alignment = top_plain
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"

    import io as _io
    stream = _io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
